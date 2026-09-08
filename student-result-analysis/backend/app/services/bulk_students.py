from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from pydantic import TypeAdapter, ValidationError
from pydantic.networks import EmailStr
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models import Login, Student
from app.security import hash_password

MAX_FILE_BYTES = 10 * 1024 * 1024
USN_RE = re.compile(r"^\d[A-Z]{2}\d{2}[A-Z]{2,4}\d{3}$")
EMAIL_ADAPTER = TypeAdapter(EmailStr)

HEADER_ALIASES = {
    "usn": {"usn", "usnno", "usnnumber", "studentusn", "studentid", "studentnumber"},
    "name": {"studentname", "fullname", "name", "studentfullname"},
    "email": {"email", "emailid", "studentemail"},
    "department": {"department", "dept", "branch", "stream"},
    "semester": {"semester", "sem"},
    "gender": {"gender", "sex"},
    "password": {"initialpassword", "password", "studentpassword", "defaultpassword"},
}
REQUIRED_COLUMNS = ("usn", "name", "email", "department", "semester", "password")
COLUMN_LABELS = {
    "usn": "USN",
    "name": "Student Name",
    "email": "Email",
    "department": "Department",
    "semester": "Semester",
    "password": "Initial Password",
}


@dataclass
class BulkRowError:
    row: int
    usn: str | None
    error: str


@dataclass
class ParsedBulkStudent:
    row: int
    usn: str
    name: str
    email: str
    department: str
    semester: int
    gender: str | None
    password: str


@dataclass
class ParsedBulkWorkbook:
    sheet_name: str
    students: list[ParsedBulkStudent] = field(default_factory=list)
    errors: list[BulkRowError] = field(default_factory=list)
    duplicate_usns: list[str] = field(default_factory=list)
    missing_required_columns: list[str] = field(default_factory=list)


class BulkImportValidationError(Exception):
    def __init__(self, message: str, parsed: ParsedBulkWorkbook):
        super().__init__(message)
        self.message = message
        self.parsed = parsed


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", _text(value).casefold())


def _header_field(value: Any) -> str | None:
    key = _header_key(value)
    aliases = {alias: field for field, names in HEADER_ALIASES.items() for alias in names}
    if key in aliases:
        return aliases[key]
    if "usn" in key or key.startswith("studentid"):
        return "usn"
    if "email" in key:
        return "email"
    if "department" in key or key in {"dept", "branch", "stream"}:
        return "department"
    if "semester" in key or key.startswith("sem"):
        return "semester"
    if "password" in key:
        return "password"
    if "gender" in key or key == "sex":
        return "gender"
    if "name" in key and ("student" in key or "full" in key or key == "name"):
        return "name"
    return None


def _find_header_row(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]] | None:
    for index, row in enumerate(rows[:50]):
        columns: dict[str, int] = {}
        for col, value in enumerate(row):
            field = _header_field(value)
            if field is not None and field not in columns:
                columns[field] = col
        if "usn" in columns and "name" in columns:
            return index, columns
    return None


def parse_student_workbook(content: bytes, filename: str) -> ParsedBulkWorkbook:
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("File is larger than 10MB")
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("Upload an Excel workbook (.xlsx or .xlsm)")
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError("Could not read the Excel file. Confirm it is a valid workbook.") from exc
    if not workbook.sheetnames:
        raise ValueError("The Excel workbook has no sheets")

    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or not any(_text(value) for row in rows for value in row):
        raise ValueError("The Excel sheet is empty")
    header = _find_header_row(rows)
    if header is None:
        raise ValueError("Could not find a header row containing USN and student name")

    header_index, columns = header
    parsed = ParsedBulkWorkbook(sheet_name=sheet.title)
    parsed.missing_required_columns = [
        COLUMN_LABELS[column]
        for column in REQUIRED_COLUMNS
        if column not in columns
    ]
    if parsed.missing_required_columns:
        return parsed

    seen: dict[str, int] = {}
    seen_emails: dict[str, int] = {}
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values = {field: _text(row[col] if col < len(row) else None) for field, col in columns.items()}
        if not any(values.values()):
            continue
        usn = values["usn"].upper()
        row_errors: list[str] = []
        if not usn:
            row_errors.append("USN is required")
        elif not USN_RE.fullmatch(usn):
            row_errors.append("USN must match the format 4MH24MC001")
        if usn in seen:
            parsed.duplicate_usns.append(usn)
            row_errors.append(f"Duplicate USN; first appears on row {seen[usn]}")
        elif usn:
            seen[usn] = row_number

        name = values["name"]
        if not name:
            row_errors.append("Student name is required")
        elif len(name) > 100:
            row_errors.append("Student name must be 100 characters or fewer")
        email = values["email"].casefold()
        if not email:
            row_errors.append("Email is required")
        else:
            try:
                email = str(EMAIL_ADAPTER.validate_python(email))
            except ValidationError:
                row_errors.append("Email must be valid")
            if len(email) > 100:
                row_errors.append("Email must be 100 characters or fewer")
            if email in seen_emails:
                row_errors.append(f"Duplicate email; first appears on row {seen_emails[email]}")
            else:
                seen_emails[email] = row_number
        department = values["department"]
        if not department:
            row_errors.append("Department is required")
        elif len(department) > 80:
            row_errors.append("Department must be 80 characters or fewer")
        password = values["password"]
        if not password:
            row_errors.append("Initial Password is required")
        elif len(password) < 6 or len(password) > 64:
            row_errors.append("Initial Password must be between 6 and 64 characters")
        try:
            semester = int(float(values["semester"]))
            if not 1 <= semester <= 8:
                raise ValueError
        except ValueError:
            semester = 0
            row_errors.append("Semester must be an integer from 1 to 8")
        gender = values.get("gender") or None
        if gender and len(gender) > 10:
            row_errors.append("Gender must be 10 characters or fewer")

        if row_errors:
            parsed.errors.append(BulkRowError(row_number, usn or None, "; ".join(row_errors)))
            continue
        parsed.students.append(
            ParsedBulkStudent(row_number, usn, name, email, department, semester, gender, password)
        )

    parsed.duplicate_usns = sorted(set(parsed.duplicate_usns))
    if not parsed.students and not parsed.errors:
        parsed.errors.append(BulkRowError(header_index + 2, None, "The workbook contains no student rows"))
    return parsed


def persist_bulk_students(db: Session, parsed: ParsedBulkWorkbook) -> int:
    if parsed.missing_required_columns or parsed.errors or parsed.duplicate_usns:
        raise BulkImportValidationError("Bulk student import was rejected", parsed)

    usns = [student.usn for student in parsed.students]
    emails = [student.email for student in parsed.students]
    existing_usns = {
        value
        for value, in db.query(Student.usn).filter(Student.usn.in_(usns)).all()
    }
    existing_login_usns = {
        value
        for value, in db.query(Login.usn).filter(Login.usn.in_(usns)).all()
    }
    existing_emails = {
        value.casefold()
        for value, in db.query(Login.email).filter(Login.email.in_(emails)).all()
    }
    for student in parsed.students:
        conflicts: list[str] = []
        if student.usn in existing_usns or student.usn in existing_login_usns:
            conflicts.append("USN already exists")
            parsed.duplicate_usns.append(student.usn)
        if student.email.casefold() in existing_emails:
            conflicts.append("Email already exists")
        if conflicts:
            parsed.errors.append(BulkRowError(student.row, student.usn, "; ".join(conflicts)))
    parsed.duplicate_usns = sorted(set(parsed.duplicate_usns))
    if parsed.errors:
        raise BulkImportValidationError("Bulk student import was rejected", parsed)

    try:
        for student in parsed.students:
            login = Login(
                usn=student.usn,
                email=student.email,
                password=hash_password(student.password),
                role="student",
            )
            db.add(login)
            db.flush()
            db.add(
                Student(
                    login_id=login.login_id,
                    usn=student.usn,
                    student_name=student.name,
                    gender=student.gender,
                    department=student.department,
                    semester=student.semester,
                )
            )
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise BulkImportValidationError(
            "Bulk import was rolled back because a student or email already exists",
            parsed,
        ) from exc
    return len(parsed.students)